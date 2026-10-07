"""Excel formatting and atomic saving for an already generated ticket report."""
from datetime import datetime, timedelta, timezone
import os
from pathlib import Path
from tempfile import NamedTemporaryFile


# Field, heading, Excel width. The GUI uses these same fields and headings.
REPORT_COLUMNS = (
    ('ticket_id', 'Ticket ID', 12),
    ('employee_name', 'Employee Name', 24),
    ('department', 'Department', 22),
    ('category', 'Category', 18),
    ('subject', 'Subject', 45),
    ('priority', 'Priority', 14),
    ('status', 'Status', 20),
    ('assigned_to', 'Assigned Technician', 30),
    ('created_at', 'Created At', 24),
    ('updated_at', 'Updated At', 24),
    ('resolved_at', 'Resolved At', 24),
)
FILTER_LABELS = ('Status', 'Priority', 'Category', 'Assigned Technician')
MANILA_TIMEZONE = timezone(timedelta(hours=8), 'Asia/Manila')
HEADER_ROW = 10


class ReportExportError(Exception):
    """Friendly feedback for dependency, filename, or filesystem failures."""


def report_generated_at():
    return datetime.now(MANILA_TIMEZONE)


def report_row_values(ticket):
    values = []
    for field, _, _ in REPORT_COLUMNS:
        value = ticket.get(field)
        if value is None or value == '':
            value = 'Unassigned' if field == 'assigned_to' else '-'
        elif isinstance(value, datetime):
            value = value.strftime('%Y-%m-%d %H:%M:%S')
        values.append(''.join(character if character.isprintable() else repr(character)[1:-1]
                              for character in str(value)))
    return tuple(values)


def _set_text(cell, value):
    # Store user text as text, including leading '='; never create a formula.
    # Replace controls that are illegal in XML with visible escapes.
    cell.value = ''.join(character if character >= ' ' or character in '\n\r\t'
                         else repr(character)[1:-1] for character in str(value))
    cell.data_type = 's'


def build_report_workbook(rows, filters, generated_at=None):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ReportExportError('Excel export requires openpyxl. Run: '
                                '.venv/Scripts/python.exe -m pip install -r requirements.txt') from None
    generated_at = generated_at if generated_at is not None else report_generated_at()
    if generated_at.tzinfo is not None:
        generated_at = generated_at.astimezone(MANILA_TIMEZONE)
    workbook = Workbook()
    workbook.properties.creator = 'IT Help Desk Ticketing System'
    sheet = workbook.active
    sheet.title = 'Ticket Report'
    last_column = get_column_letter(len(REPORT_COLUMNS))
    sheet.merge_cells(f'A1:{last_column}1')
    _set_text(sheet['A1'], 'IT Help Desk Ticket Report')
    sheet['A1'].font = Font(size=18, bold=True, color='182A43')
    sheet.row_dimensions[1].height = 30
    sheet.merge_cells(f'A2:{last_column}2')
    _set_text(sheet['A2'], f'Generated: {generated_at:%Y-%m-%d %H:%M:%S} Asia/Manila (UTC+08:00)')
    for row_number, label in enumerate(FILTER_LABELS, start=4):
        _set_text(sheet.cell(row_number, 1), label)
        sheet.cell(row_number, 1).font = Font(bold=True)
        sheet.merge_cells(start_row=row_number, start_column=2, end_row=row_number, end_column=len(REPORT_COLUMNS))
        _set_text(sheet.cell(row_number, 2), filters.get(label, 'All'))
    _set_text(sheet['A8'], f'{len(rows)} tickets found.')
    for column_number, (field, heading, width) in enumerate(REPORT_COLUMNS, start=1):
        cell = sheet.cell(HEADER_ROW, column_number)
        _set_text(cell, heading)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='182A43')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        sheet.column_dimensions[get_column_letter(column_number)].width = width
    sheet.row_dimensions[HEADER_ROW].height = 30
    for row_number, ticket in enumerate(rows, start=HEADER_ROW + 1):
        for column_number, (field, _, _) in enumerate(REPORT_COLUMNS, start=1):
            cell = sheet.cell(row_number, column_number)
            value = ticket.get(field)
            if value is None or value == '':
                if field == 'assigned_to':
                    _set_text(cell, 'Unassigned')
            elif isinstance(value, datetime):
                cell.value = (value.astimezone(MANILA_TIMEZONE).replace(tzinfo=None)
                              if value.tzinfo is not None else value)
                cell.number_format = 'yyyy-mm-dd hh:mm:ss'
            elif field == 'ticket_id' and type(value) is int:
                cell.value = value
            else:
                _set_text(cell, value)
            cell.alignment = Alignment(vertical='top', wrap_text=field in ('employee_name', 'subject', 'assigned_to'))
            if row_number % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F0F4FA')
    sheet.freeze_panes = f'A{HEADER_ROW + 1}'
    sheet.auto_filter.ref = f'A{HEADER_ROW}:{last_column}{HEADER_ROW + len(rows)}'
    return workbook


def export_ticket_report(path, rows, filters, generated_at=None):
    """Replace the chosen file only after the complete workbook saves successfully."""
    path = Path(path).absolute()
    if path.suffix.lower() != '.xlsx':
        raise ReportExportError('Choose a filename ending in .xlsx.')
    workbook, temporary_path = None, None
    try:
        workbook = build_report_workbook(rows, filters, generated_at)
        with NamedTemporaryFile(dir=path.parent, prefix='.helpdesk_report_', suffix='.xlsx', delete=False) as temporary:
            temporary_path = Path(temporary.name)
        workbook.save(temporary_path)
        os.replace(temporary_path, path)
        return str(path)
    except OSError:
        raise ReportExportError('Unable to save the Excel report. Choose a writable location '
                                'and close the destination file in Excel before retrying.') from None
    finally:
        if workbook is not None:
            workbook.close()
        if temporary_path is not None:
            try:
                temporary_path.unlink(missing_ok=True)
            except OSError:
                pass  # Do not hide the original save error if cleanup is also denied.
