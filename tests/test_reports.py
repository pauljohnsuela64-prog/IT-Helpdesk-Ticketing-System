"""Read-only report queries, permissions, and actual Excel round trips."""
from contextlib import closing
from datetime import datetime, timezone
from pathlib import Path
import sqlite3
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import MagicMock, patch

import mysql.connector
from openpyxl import load_workbook

import report_export as export
import report_repository as repo


def report_ticket(ticket_id=3, **overrides):
    return dict(ticket_id=ticket_id, employee_name='Alice Reyes', department='IT', category='Hardware',
                subject='PC 3 issue', priority='Medium', status='Open', assigned_to=None,
                created_at=datetime(2026, 10, 7, 9, 0), updated_at=datetime(2026, 10, 7, 9, 5),
                resolved_at=None, **overrides)


class ReportRepositoryTests(unittest.TestCase):
    def setUp(self):
        self.connection = MagicMock()
        self.connection.__enter__.return_value = self.connection
        self.cursor = self.connection.cursor.return_value.__enter__.return_value
        self.cursor.fetchone.return_value = dict(user_id=7, role='Admin', status='Active')
        self.cursor.fetchall.return_value = [report_ticket()]
        self.connect = patch.object(repo, 'get_connection', return_value=self.connection).start()
        self.addCleanup(patch.stopall)

    def test_all_filters_select_only_report_columns_in_stable_order(self):
        self.assertEqual(repo.get_ticket_report(7), [report_ticket()])
        query, parameters = self.cursor.execute.call_args.args
        self.assertNotIn('WHERE', query)
        self.assertIn('FROM helpdesk.tickets ORDER BY ticket_id', query)
        self.assertEqual(parameters, ())
        for field, _, _ in export.REPORT_COLUMNS:
            self.assertIn(field, query)
        self.assertNotIn('password', query.lower())
        self.assertNotIn('description', query)
        self.connection.commit.assert_not_called()

    def test_optional_filters_are_bound_and_combined_with_and(self):
        repo.get_ticket_report(7, status='Resolved', priority='Critical', category='Network', assigned_technician_id=12)
        query, parameters = self.cursor.execute.call_args.args
        self.assertIn('WHERE status = %s AND priority = %s AND category = %s AND assigned_technician_id = %s', query)
        self.assertEqual(parameters, ('Resolved', 'Critical', 'Network', 12))
        for value in parameters[:-1]:
            self.assertNotIn(value, query)
        self.assertEqual(query.count('%s'), len(parameters))

    def test_each_single_filter_is_optional(self):
        for filters, expected in (({'status': 'Open'}, ('Open',)), ({'priority': 'High'}, ('High',)),
                                  ({'category': 'Email'}, ('Email',)), ({'assigned_technician_id': 3}, (3,)),
                                  ({'status': None, 'priority': None, 'category': None}, ())):
            with self.subTest(filters=filters):
                repo.get_ticket_report(7, **filters)
                self.assertEqual(self.cursor.execute.call_args.args[1], expected)

    def test_invalid_and_injection_filters_are_rejected_before_sql(self):
        for filters in ({'status': "Open' OR 1=1 --"}, {'priority': 'Urgent'}, {'category': 'Unknown'},
                        {'assigned_technician_id': '3 OR 1=1'}, {'assigned_technician_id': True},
                        {'assigned_technician_id': 0}):
            with self.subTest(filters=filters), self.assertRaises(ValueError):
                repo.get_ticket_report(7, **filters)
        self.connect.assert_not_called()

    def test_every_read_and_export_authorization_rechecks_active_admin_without_passwords(self):
        for operation in (repo.get_ticket_report, repo.get_report_technicians, repo.authorize_report_access):
            with self.subTest(operation=operation.__name__):
                self.cursor.execute.reset_mock()
                operation(7)
                self.assertEqual(self.cursor.execute.call_args_list[0].args,
                                 ('SELECT user_id, role, status FROM helpdesk.users WHERE user_id = %s LIMIT 1', (7,)))
                self.assertTrue(all(call.args[0].startswith('SELECT ') for call in self.cursor.execute.call_args_list))
                self.assertTrue(all('password' not in call.args[0].lower() for call in self.cursor.execute.call_args_list))
                self.connection.commit.assert_not_called()

    def test_technician_inactive_missing_and_unknown_accounts_cannot_read_or_export(self):
        for actor in (None, dict(role='Technician', status='Active'), dict(role='Admin', status='Inactive'),
                      dict(role='Unknown', status='Active')):
            for operation in (repo.get_ticket_report, repo.get_report_technicians, repo.authorize_report_access):
                with self.subTest(actor=actor, operation=operation.__name__):
                    self.cursor.execute.reset_mock()
                    self.cursor.fetchone.return_value = actor
                    with self.assertRaisesRegex(repo.ReportPermissionError, 'permission'):
                        operation(7)
                    self.assertEqual(self.cursor.execute.call_count, 1)

    def test_invalid_user_id_cannot_open_database_or_fall_back_to_public_reports(self):
        for user_id in (None, True, 0, '7', -1):
            for operation in (repo.get_ticket_report, repo.get_report_technicians, repo.authorize_report_access):
                with self.subTest(user_id=user_id), self.assertRaises(repo.ReportPermissionError):
                    operation(user_id)
        self.connect.assert_not_called()

    def test_technician_options_include_inactive_records_and_no_authentication_fields(self):
        rows = [dict(technician_id=12, full_name='Mark Santos', status='Inactive')]
        self.cursor.fetchall.return_value = rows
        self.assertEqual(repo.get_report_technicians(7), rows)
        query, parameters = self.cursor.execute.call_args.args
        self.assertEqual(query, 'SELECT technician_id, full_name, status FROM helpdesk.technicians ORDER BY technician_id')
        self.assertEqual(parameters, ())

    def test_empty_and_database_failures_are_friendly_without_private_details(self):
        self.cursor.fetchall.return_value = []
        self.assertEqual(repo.get_ticket_report(7, status='Closed'), [])
        for error in (ValueError('private settings'), mysql.connector.Error('private detail', errno=2003)):
            self.connect.side_effect = error
            with self.assertRaises(repo.ReportReadError) as raised:
                repo.get_ticket_report(7)
            self.assertNotIn('private', str(raised.exception))

    def test_actual_filter_query_uses_assignment_ids_for_duplicate_names(self):
        with closing(sqlite3.connect(':memory:')) as db:
            db.execute("ATTACH DATABASE ':memory:' AS helpdesk")
            db.execute('CREATE TABLE helpdesk.tickets (ticket_id INT, employee_name TEXT, department TEXT, '
                       'category TEXT, subject TEXT, priority TEXT, status TEXT, assigned_to TEXT, '
                       'created_at TEXT, updated_at TEXT, resolved_at TEXT, assigned_technician_id INT)')
            db.executemany('INSERT INTO helpdesk.tickets VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)', [
                (3, 'Alice', 'IT', 'Network', 'Cable fault', 'High', 'Resolved', 'Same Name', None, None, None, 12),
                (4, 'Bob', 'IT', 'Network', 'Cable fault', 'High', 'Resolved', 'Same Name', None, None, None, 35),
                (5, 'Anna', 'IT', 'Hardware', 'PC issue', 'High', 'Open', None, None, None, None, None),
            ])
            db.row_factory = sqlite3.Row
            result = []

            def execute(query, parameters):
                if 'FROM helpdesk.users' not in query:
                    result[:] = [dict(row) for row in db.execute(query.replace('%s', '?'), parameters)]

            self.cursor.execute.side_effect = execute
            self.cursor.fetchall.side_effect = lambda: result
            self.assertEqual([row['ticket_id'] for row in repo.get_ticket_report(7)], [3, 4, 5])
            matching = repo.get_ticket_report(7, status='Resolved', priority='High', category='Network', assigned_technician_id=12)
            self.assertEqual([row['ticket_id'] for row in matching], [3])
            self.assertEqual(repo.get_ticket_report(7, status='Closed'), [])


class ExcelReportTests(unittest.TestCase):
    def setUp(self):
        self.directory = TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / 'ticket report.xlsx'
        self.filters = dict(zip(export.FILTER_LABELS, ('All', 'High', 'Hardware', 'Mark Santos (ID: 12)')))
        self.generated_at = datetime(2026, 10, 7, 1, 15, tzinfo=timezone.utc)

    def open_export(self, rows, filters=None):
        saved = export.export_ticket_report(self.path, rows, self.filters if filters is None else filters, self.generated_at)
        self.assertEqual(saved, str(self.path))
        workbook = load_workbook(self.path)
        self.addCleanup(workbook.close)
        return workbook.active

    def test_round_trip_contains_title_local_time_filters_headers_and_all_rows(self):
        rows = [report_ticket(), {**report_ticket(9), 'employee_name': 'José Reyes', 'assigned_to': 'Mark Santos',
                                  'resolved_at': datetime(2026, 10, 7, 10, 30)}]
        sheet = self.open_export(rows)
        self.assertEqual(sheet['A1'].value, 'IT Help Desk Ticket Report')
        self.assertEqual(sheet['A2'].value, 'Generated: 2026-10-07 09:15:00 Asia/Manila (UTC+08:00)')
        for row_number, label in enumerate(export.FILTER_LABELS, start=4):
            self.assertEqual(sheet.cell(row_number, 1).value, label)
            self.assertEqual(sheet.cell(row_number, 2).value, self.filters[label])
        self.assertEqual(sheet['A8'].value, '2 tickets found.')
        self.assertEqual(tuple(cell.value for cell in sheet[10]), tuple(heading for _, heading, _ in export.REPORT_COLUMNS))
        self.assertEqual([sheet.cell(row, 1).value for row in (11, 12)], [3, 9])
        self.assertEqual(sheet['B12'].value, 'José Reyes')
        self.assertEqual(sheet['H11'].value, 'Unassigned')
        self.assertEqual(sheet['H12'].value, 'Mark Santos')
        self.assertEqual(sheet['I11'].value, rows[0]['created_at'])
        self.assertIsNone(sheet['K11'].value)
        self.assertEqual(sheet['K12'].value, rows[1]['resolved_at'])
        self.assertEqual(sheet.max_row, 12)

    def test_formatting_preserves_numeric_ids_dates_and_readability(self):
        sheet = self.open_export([report_ticket()])
        self.assertTrue(all(cell.font.bold for cell in sheet[10]))
        self.assertEqual(sheet.freeze_panes, 'A11')
        self.assertEqual(sheet.auto_filter.ref, 'A10:K11')
        self.assertGreaterEqual(sheet.column_dimensions['E'].width, 40)
        self.assertTrue(sheet['E11'].alignment.wrap_text)
        self.assertEqual(sheet['A11'].data_type, 'n')
        self.assertEqual(sheet['I11'].number_format, 'yyyy-mm-dd hh:mm:ss')

    def test_user_text_never_becomes_formulas_hyperlinks_or_excel_errors(self):
        for value in ('=HYPERLINK("https://example.invalid", "Click")', '=1+1', '#REF!', '+123', '-123', '@SUM(A1)',
                      'First line\nSecond line\x00'):
            with self.subTest(value=value):
                sheet = self.open_export([{**report_ticket(), 'subject': value}])
                self.assertEqual(sheet['E11'].data_type, 's')
                self.assertIsNone(sheet['E11'].hyperlink)
                self.assertEqual(sheet['E11'].value, value.replace('\x00', '\\x00'))

    def test_authentication_and_unrelated_fields_are_never_exported(self):
        row = {**report_ticket(), 'password': 'excluded_sensitive_field', 'password_hash': 'excluded_sensitive_field',
               'username': 'excluded_account_identity', 'description': 'excluded_unrequested_description'}
        sheet = self.open_export([row])
        exported = str(tuple(sheet.values))
        for value in ('excluded_sensitive_field', 'excluded_account_identity', 'excluded_unrequested_description'):
            self.assertNotIn(value, exported)

    def test_empty_report_still_has_metadata_and_header(self):
        sheet = self.open_export([], {})
        self.assertEqual(sheet['A8'].value, '0 tickets found.')
        self.assertEqual(sheet.max_row, 10)
        self.assertEqual(sheet.auto_filter.ref, 'A10:K10')
        self.assertEqual(sheet['B4'].value, 'All')

    def test_invalid_extension_and_missing_directory_are_friendly(self):
        for path in (self.path.with_suffix('.csv'), self.path.parent / 'missing' / 'report.xlsx'):
            with self.subTest(path=path), self.assertRaises(export.ReportExportError):
                export.export_ticket_report(path, [report_ticket()], {})
        self.assertFalse(self.path.exists())

    def test_failed_save_or_replace_keeps_existing_file_and_cleans_temporary_file(self):
        self.path.write_bytes(b'previous report')
        for target in ('openpyxl.workbook.workbook.Workbook.save', 'report_export.os.replace'):
            with self.subTest(target=target), patch(target, side_effect=PermissionError('private filesystem detail')):
                with self.assertRaises(export.ReportExportError) as raised:
                    export.export_ticket_report(self.path, [report_ticket()], {})
                self.assertNotIn('private', str(raised.exception))
            self.assertEqual(self.path.read_bytes(), b'previous report')
            self.assertEqual(list(self.path.parent.iterdir()), [self.path])

    def test_missing_dependency_is_friendly_without_affecting_gui_imports(self):
        with patch.dict('sys.modules', {'openpyxl': None}):
            with self.assertRaisesRegex(export.ReportExportError, 'pip install -r requirements.txt'):
                export.build_report_workbook([], {})

    def test_preview_formats_datetimes_nulls_and_multiline_text_without_mutating_data(self):
        row = report_ticket()
        row['subject'] = 'PC\n3\tissue'
        original = row.copy()
        values = export.report_row_values(row)
        self.assertEqual(values[0], '3')
        self.assertEqual(values[4], 'PC\\n3\\tissue')
        self.assertEqual(values[7], 'Unassigned')
        self.assertEqual(values[8], '2026-10-07 09:00:00')
        self.assertEqual(values[10], '-')
        self.assertEqual(row, original)


if __name__ == '__main__':
    unittest.main()
